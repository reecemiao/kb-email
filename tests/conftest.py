import struct
import zlib
from email.message import EmailMessage
from email.utils import format_datetime, make_msgid
from datetime import datetime, timedelta

import pytest
import yaml

from kbmail.workspace import Workspace, init_workspace


def png_bytes(size: int = 9000, seed: int = 0) -> bytes:
    """A valid PNG padded with a text chunk so it passes the min-size filter."""
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    idat = zlib.compress(b"\x00" + bytes([seed % 256, 0, 0]))
    pad = chunk(b"tEXt", b"pad\x00" + bytes(size))
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + pad + chunk(b"IDAT", idat) + chunk(b"IEND", b"")


def docx_bytes(text: str) -> bytes:
    import io
    import docx  # python-docx
    d = docx.Document()
    d.add_heading("Spec", 1)
    d.add_paragraph(text)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "key", "value", "timeout", "30s"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


SVG = """<svg xmlns="http://www.w3.org/2000/svg" width="300" height="100">
<rect x="10" y="10" width="80" height="40"/><text x="20" y="35">Client</text>
<rect x="200" y="10" width="80" height="40"/><text x="210" y="35">API</text>
<line x1="90" y1="30" x2="200" y2="30"/></svg>"""


def make_eml(path, *, subject, frm, to, body, cc=(), date=None, msgid=None, refs=None, html=None,
             inline_png=None, attachments=()):
    m = EmailMessage()
    m["Subject"] = subject
    m["From"] = frm
    m["To"] = ", ".join(to)
    if cc:
        m["Cc"] = ", ".join(cc)
    m["Date"] = format_datetime(date or datetime.now().astimezone())
    m["Message-ID"] = msgid or make_msgid()
    if refs:
        m["References"] = refs
        m["In-Reply-To"] = refs.split()[-1]
    m.set_content(body)
    if html:
        m.add_alternative(html, subtype="html")
        if inline_png:
            m.get_payload()[1].add_related(inline_png, "image", "png", cid="<img001@x>")
    for name, data, maintype, subtype in attachments:
        m.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
    path.write_bytes(bytes(m))
    return m["Message-ID"]


@pytest.fixture
def ws(tmp_path):
    root = tmp_path / "kbws"
    init_workspace(root)
    cfg = yaml.safe_load((root / "kbmail.yaml").read_text(encoding="utf-8"))
    cfg["me"] = {"name": "Taro Yamada", "emails": ["taro@contoso.com"]}
    cfg["areas"] = [
        {"name": "payments", "description": "Payment API and billing", "keywords": ["payment", "billing"]},
        {"name": "infra", "description": "Servers, network, cloud", "keywords": ["server", "network"]},
    ]
    (root / "kbmail.yaml").write_text(yaml.safe_dump(cfg, allow_unicode=True), encoding="utf-8")
    return Workspace(root)


@pytest.fixture
def mailbox(tmp_path):
    return build_mailbox(tmp_path / "eml")


def build_mailbox(d):
    """A small mailbox: a payments thread with attachments, infra mail, a newsletter, a Japanese mail."""
    d.mkdir(parents=True, exist_ok=True)
    now = datetime.now().astimezone().replace(microsecond=0)
    root_id = make_eml(
        d / "01.eml", subject="Payment API cut-over plan", frm="Sarah Tanaka <sarah@contoso.com>",
        to=["taro@contoso.com"], cc=["ken@vendor.example.com"], date=now - timedelta(days=5),
        body="Hi Taro,\nThe cut-over is planned for 2026-11-15. Could you confirm the rollback criteria?\n"
             "password: hunter2secret\nThanks, Sarah",
        html="<html><body><p>Hi Taro,</p><p>The cut-over is planned for <b>2026-11-15</b>.</p>"
             "<table><tr><th>step</th><th>owner</th></tr><tr><td>freeze</td><td>Sarah</td></tr></table>"
             "<p><img src=\"cid:img001@x\"></p><p>password: hunter2secret</p></body></html>",
        inline_png=png_bytes(9000, 1),
        attachments=[("spec.docx", docx_bytes("Payment API v2 uses idempotency keys."), "application",
                      "vnd.openxmlformats-officedocument.wordprocessingml.document"),
                     ("flow.svg", SVG.encode(), "image", "svg+xml")],
    )
    make_eml(d / "02.eml", subject="RE: Payment API cut-over plan", frm="Taro Yamada <taro@contoso.com>",
             to=["sarah@contoso.com"], date=now - timedelta(days=4), refs=root_id,
             body="Sarah,\nRollback if error rate > 2% for 10 minutes.\n\nFrom: Sarah Tanaka\nSent: Monday\n"
                  "To: Taro\nSubject: Payment API cut-over plan\n\nHi Taro, old quoted text")
    make_eml(d / "03.eml", subject="Server patching window", frm="Ops Bot <ops@contoso.com>",
             to=["taro@contoso.com"], date=now - timedelta(days=3),
             body="Servers in rack 4 will be patched on Saturday. Network downtime 10 minutes.")
    make_eml(d / "04.eml", subject="Weekly newsletter", frm="News <noreply@news.example.org>",
             to=["taro@contoso.com"], date=now - timedelta(days=2), body="Top stories this week ...")
    make_eml(d / "05.eml", subject="請求システムの件", frm="佐藤 花子 <hanako@contoso.com>",
             to=["taro@contoso.com"], date=now - timedelta(days=1),
             body="山田さん\n請求システムの移行について、ご確認をお願いします。\nよろしくお願いいたします。")
    return d
